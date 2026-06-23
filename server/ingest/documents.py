"""Document text extraction: pdf, txt, md, csv, html, docx, pptx.

Heavy/optional parsers are imported lazily so the core install stays light.
Each extractor returns a list of TextPiece with a provenance locator
(e.g. page number) where available.
"""
from __future__ import annotations

import csv
import html
import re
from pathlib import Path

from ..chunking import TextPiece

_TAG_RE = re.compile(r"<[^>]+>")
_WS_RE = re.compile(r"[ \t]+")


def _clean(text: str) -> str:
    return _WS_RE.sub(" ", text).strip()


def extract_txt(path: Path) -> list[TextPiece]:
    text = path.read_text(encoding="utf-8", errors="replace")
    return [TextPiece(text=text)] if text.strip() else []


def extract_html(path: Path) -> list[TextPiece]:
    raw = path.read_text(encoding="utf-8", errors="replace")
    raw = re.sub(r"<(script|style)[^>]*>.*?</\1>", " ", raw, flags=re.S | re.I)
    text = html.unescape(_TAG_RE.sub(" ", raw))
    text = "\n".join(_clean(line) for line in text.splitlines() if _clean(line))
    return [TextPiece(text=text)] if text.strip() else []


def extract_csv(path: Path) -> list[TextPiece]:
    pieces: list[TextPiece] = []
    with path.open("r", encoding="utf-8", errors="replace", newline="") as fh:
        reader = csv.reader(fh)
        rows = list(reader)
    if not rows:
        return []
    header = rows[0]
    body = []
    for i, row in enumerate(rows[1:], start=2):
        cells = ", ".join(f"{h}: {v}" for h, v in zip(header, row, strict=False))
        body.append(cells or ", ".join(row))
        if len(body) >= 50:  # group rows into pieces for chunking
            pieces.append(TextPiece(text="\n".join(body), locator=f"rows {i-49}-{i}"))
            body = []
    if body:
        pieces.append(TextPiece(text="\n".join(body), locator="rows (tail)"))
    return pieces


def extract_pdf(path: Path, ocr_fallback: bool = True) -> list[TextPiece]:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise RuntimeError(
            "PDF support needs 'pypdf' (in requirements.txt). Run: pip install pypdf"
        ) from exc
    reader = PdfReader(str(path))
    pieces: list[TextPiece] = []
    for i, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if not text and ocr_fallback:
            text = _ocr_pdf_page(path, i - 1)
        if text:
            pieces.append(TextPiece(text=text, locator=f"p. {i}"))
    return pieces


def _ocr_pdf_page(path: Path, page_index: int) -> str:
    """Best-effort OCR for a scanned PDF page. Optional deps; returns '' if absent."""
    try:
        import pytesseract
        from pdf2image import convert_from_path
    except ImportError:
        return ""
    try:
        images = convert_from_path(str(path), first_page=page_index + 1,
                                   last_page=page_index + 1)
        if not images:
            return ""
        return pytesseract.image_to_string(images[0]).strip()
    except Exception:
        return ""


def extract_docx(path: Path) -> list[TextPiece]:
    try:
        import docx  # python-docx
    except ImportError as exc:
        raise RuntimeError(
            "DOCX support needs 'python-docx'. Run: pip install python-docx"
        ) from exc
    doc = docx.Document(str(path))
    text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    return [TextPiece(text=text)] if text.strip() else []


def extract_pptx(path: Path) -> list[TextPiece]:
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise RuntimeError(
            "PPTX support needs 'python-pptx'. Run: pip install python-pptx"
        ) from exc
    prs = Presentation(str(path))
    pieces: list[TextPiece] = []
    for i, slide in enumerate(prs.slides, start=1):
        texts = [sh.text for sh in slide.shapes if getattr(sh, "has_text_frame", False)]
        joined = "\n".join(t for t in texts if t.strip())
        if joined.strip():
            pieces.append(TextPiece(text=joined, locator=f"slide {i}"))
    return pieces


_DISPATCH = {
    ".txt": extract_txt, ".md": extract_txt,
    ".html": extract_html, ".htm": extract_html,
    ".csv": extract_csv,
    ".pdf": extract_pdf,
    ".docx": extract_docx,
    ".pptx": extract_pptx,
}


def extract_document(path: Path) -> list[TextPiece]:
    fn = _DISPATCH.get(path.suffix.lower())
    if fn is None:
        raise RuntimeError(f"Unsupported document type: {path.suffix}")
    return fn(path)
